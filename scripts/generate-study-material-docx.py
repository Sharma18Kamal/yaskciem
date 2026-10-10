from __future__ import annotations

import argparse
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "user-guide" / "study-material.md"
OUTPUT = ROOT / "docs" / "user-guide" / "YASK-CIEM-Study-Material.docx"
IMAGE_DIR = ROOT / "docs" / "user-guide" / "images"

SCREENSHOTS = {
    "Part 4 - Running a scan (Dashboard)": ("01-dashboard.png", "Dashboard overview"),
    "5.1 Identities": ("04-resources.png", "Identities page"),
    "5.3 Entitlements": ("05-entitlements.png", "Entitlements page"),
    "5.4 Effective Access": ("06-effective-access.png", "Effective Access page"),
    "5.5 Access Paths": ("07-access-paths.png", "Access Paths page"),
    "5.6 Risk Findings": ("08-risk-findings.png", "Risk Findings page"),
    "6.1 Recommendations": ("09-recommendations.png", "Recommendations page"),
    "6.2 Remediation (approval-gated)": ("10-remediation.png", "Remediation page"),
    "6.3 Access Reviews": ("12-access-reviews.png", "Access Reviews page"),
    "7.1 Organizations page": ("11-organizations.png", "Organizations page"),
    "Part 8 - Settings": ("15-settings.png", "Settings page"),
}


def add_inline(paragraph, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        if node:
            run = paragraph.add_run(str(node))
            run.bold = bold
            run.italic = italic
            if code:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        return

    if not isinstance(node, Tag):
        return

    tag = node.name.lower()
    next_bold = bold or tag in {"strong", "b"}
    next_italic = italic or tag in {"em", "i"}
    next_code = code or tag == "code"
    for child in node.children:
        add_inline(paragraph, child, next_bold, next_italic, next_code)


def add_list(document, element, depth=0):
    ordered = element.name == "ol"
    style_name = "List Number" if ordered else "List Bullet"
    if depth:
        style_name = f"{style_name} {min(depth + 1, 3)}"

    for item in element.find_all("li", recursive=False):
        paragraph = document.add_paragraph(style=style_name)
        for child in item.children:
            if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                add_list(document, child, depth + 1)
            elif isinstance(child, Tag) and child.name == "p":
                for inline in child.children:
                    add_inline(paragraph, inline)
            else:
                add_inline(paragraph, child)


def add_table(document, element):
    rows = element.find_all("tr")
    if not rows:
        return
    column_count = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    table.autofit = True

    for row_index, row in enumerate(rows):
        cells = row.find_all(["th", "td"], recursive=False)
        target_cells = table.add_row().cells
        for column_index, cell in enumerate(cells):
            target = target_cells[column_index]
            target.text = ""
            paragraph = target.paragraphs[0]
            for inline in cell.children:
                add_inline(paragraph, inline, bold=(row_index == 0))
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_screenshot(document, image_name, caption):
    image_path = IMAGE_DIR / image_name
    if not image_path.is_file():
        raise FileNotFoundError(f"Missing screenshot: {image_path}")

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.9))

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(10)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(9)
    caption_run.font.color.rgb = RGBColor(89, 89, 89)


def add_element(document, element):
    tag = element.name.lower()

    if tag in {"h2", "h3", "h4"}:
        level = {"h2": 1, "h3": 2, "h4": 3}[tag]
        heading = element.get_text(" ", strip=True)
        paragraph = document.add_heading(heading, level=level)
        paragraph.paragraph_format.keep_with_next = True
        screenshot = SCREENSHOTS.get(heading)
        if screenshot:
            add_screenshot(document, *screenshot)
    elif tag == "p":
        paragraph = document.add_paragraph()
        for child in element.children:
            add_inline(paragraph, child)
    elif tag in {"ul", "ol"}:
        add_list(document, element)
    elif tag == "table":
        add_table(document, element)
        document.add_paragraph()
    elif tag == "blockquote":
        for child in element.find_all(["p"], recursive=False):
            paragraph = document.add_paragraph(style="Quote")
            for inline in child.children:
                add_inline(paragraph, inline)


def build_document(source: Path, output: Path):
    source_text = source.read_text(encoding="utf-8")
    html = markdown.markdown(source_text, extensions=["tables", "sane_lists"])
    soup = BeautifulSoup(html, "html.parser")
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("YASK-CIEM Study Material")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A visual guide to cloud identity exposure management").italic = True

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    note.add_run(
        "Screenshots are illustrative interface captures. They may show demo data "
        "and an earlier application version; labels and values can differ from your installation."
    )

    document.add_page_break()

    for element in soup.contents:
        if not isinstance(element, Tag):
            continue
        if element.name == "h1":
            continue
        if element.name == "p" and not document.paragraphs[-1].text and len(document.paragraphs) == 1:
            continue
        add_element(document, element)

    document.core_properties.title = "YASK-CIEM Study Material"
    document.core_properties.subject = "Visual user guide and study material"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser(description="Convert the YASK-CIEM study guide to Word.")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    arguments = parser.parse_args()
    build_document(arguments.source.resolve(), arguments.output.resolve())
    print(f"Created {arguments.output.resolve()}")


if __name__ == "__main__":
    main()
